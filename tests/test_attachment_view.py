"""Every attachment kind can be opened in the UI, and none of them can run.

The promise is *"any file that arrives as an attachment can be looked at without leaving the
browser"*. Before the viewer, `mail_view._preview` covered image, pdf, xlsx/xls and docx and
returned `""` for the other ten — so a `.msg` carrying the POD evidence, an HTML delivery
confirmation and a CSV tracker each rendered as a filename and a Download link.

Enforced structurally, the same way `test_attachment_universality` enforces the extraction side:
`test_every_kind_renders_a_panel` walks `sniff.ALL_KINDS`, so a sixteenth kind cannot be added
without deciding what shows it.

The other half is that "openable" never becomes "executable". Attachments come from outside
parties. Anything we can read is re-emitted as our own escaped markup; anything we cannot is
sandboxed without `allow-scripts`. Both are asserted below, and `test_operations_readonly` sweeps
the same rule across the package.
"""
import io
import zipfile

import pytest

from pipeline import attachment_view
from pipeline.parsing import sniff
from tests.test_attachment_universality import KIND_FIXTURES

HOSTILE_HTML = (b"<html><body><h1>Delivered</h1>"
                b"<script>alert('x')</script>"
                b"<img src=x onerror=alert(1)>"
                b"</body></html>")


def render(kind: str):
    filename, make = KIND_FIXTURES[kind]
    return attachment_view.render(make(), filename, "")


# --- The promise --------------------------------------------------------------

@pytest.mark.parametrize("kind", sorted(sniff.ALL_KINDS))
def test_every_kind_renders_a_panel(kind):
    """No kind may render as nothing. A blank panel is indistinguishable from a broken one, and it
    is what sent people out of the browser to find out what they had been sent."""
    panel = attachment_view.render(*_fixture(kind))
    assert panel.html.strip(), f"{kind} rendered nothing"
    assert panel.label, f"{kind} has no human label"


@pytest.mark.parametrize("kind", sorted(sniff.ALL_KINDS))
def test_no_kind_raises(kind):
    """A renderer that throws must be caught and explained, not 500 the popup."""
    panel = attachment_view.render(*_fixture(kind))
    assert isinstance(panel.html, str)


def test_a_kind_with_no_reader_says_what_it_is_and_what_to_do():
    """`.doc` has no reliable pure-Python reader. The honest panel is the correct answer — and it
    quotes `unsupported_adapter.GUIDANCE`, so the screen and the ledger say the same thing."""
    panel = render(sniff.KIND_DOC)
    assert "No reliable reader" in panel.html
    assert "First bytes:" in panel.html
    assert "Download" in panel.html


def test_the_renderer_table_covers_every_readable_kind():
    """A guard on the guard: if a kind is dropped from `_RENDERERS` it silently falls to the
    unsupported panel, which would look like a deliberate decision rather than a regression."""
    for kind in (sniff.KIND_PDF, sniff.KIND_XLSX, sniff.KIND_XLS, sniff.KIND_DOCX,
                 sniff.KIND_PPTX, sniff.KIND_HTML, sniff.KIND_TEXT, sniff.KIND_MSG,
                 sniff.KIND_IMAGE, sniff.KIND_ARCHIVE, sniff.KIND_ZIP_OFFICE):
        assert kind in attachment_view._RENDERERS, kind


# --- Safety -------------------------------------------------------------------

def test_a_hostile_html_attachment_cannot_run():
    """The attachment the raw route refuses to serve inline. Rendering it is safe only because the
    frame has no `allow-scripts` — the markup is visible, and inert."""
    panel = attachment_view.render(HOSTILE_HTML, "notice.html", "text/html")
    assert "allow-scripts" not in panel.html
    assert "sandbox=" in panel.html
    # The payload is inside `srcdoc`, escaped — never as live markup in our own document.
    assert "<script>alert" not in panel.html
    assert "&lt;script&gt;" in panel.html


def test_every_frame_this_module_emits_is_scriptless_and_same_origin():
    """Mirrors `test_operations_readonly.test_message_frame_never_allows_scripts` at the value
    level rather than the source level. With both flags a sandbox is no sandbox."""
    import re

    from pathlib import Path
    source = Path(attachment_view.__file__).read_text(encoding="utf-8")
    sandboxes = re.findall(r'sandbox="([^"]*)"', source)
    assert sandboxes, "no sandbox found — the frame helper must always sandbox"
    for value in sandboxes:
        assert "allow-scripts" not in value, value
        assert "allow-same-origin" in value, value


def test_a_spreadsheet_cell_cannot_inject_markup():
    """Cell values are attacker-controlled and are re-emitted as our own table."""
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.active.append(["<script>alert(1)</script>", '"><b>x'])
    buffer = io.BytesIO()
    workbook.save(buffer)

    panel = attachment_view.render(buffer.getvalue(), "t.xlsx", "")
    assert "<script>alert" not in panel.html
    assert "&lt;script&gt;" in panel.html


def test_a_container_member_name_cannot_inject_markup():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("<script>alert(1)</script>.txt", "hi")
    panel = attachment_view.render(buffer.getvalue(), "bundle.zip", "")
    assert "<script>alert" not in panel.html


# --- What the readable kinds actually show ------------------------------------

def test_a_spreadsheet_shows_every_sheet_not_the_first_three():
    """The old preview stopped at three sheets and sixty rows, which silently hid most of a
    tracker — and a tracker's delivery tab is rarely the first one."""
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.active.title = "Cover"
    workbook.active.append(["nothing here"])
    for name in ("Second", "Third", "Deliveries"):
        sheet = workbook.create_sheet(name)
        sheet.append(["PO", "Qty"])
        sheet.append(["208491", 11])
    buffer = io.BytesIO()
    workbook.save(buffer)

    panel = attachment_view.render(buffer.getvalue(), "tracker.xlsx", "")
    assert "Deliveries" in panel.html, "the fourth sheet was dropped"
    assert "208491" in panel.html


def test_a_document_shows_its_tables_not_only_its_paragraphs():
    """python-docx does not report table cells as paragraphs, and a delivery confirmation attached
    as .docx puts its quantities in a table — so the old paragraphs-only preview dropped exactly
    the part worth reading."""
    import docx

    document = docx.Document()
    document.add_paragraph("Delivery confirmation")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "PO"
    table.cell(0, 1).text = "Qty"
    table.cell(1, 0).text = "208491"
    table.cell(1, 1).text = "202"
    buffer = io.BytesIO()
    document.save(buffer)

    panel = attachment_view.render(buffer.getvalue(), "notice.docx", "")
    assert "Delivery confirmation" in panel.html
    assert "208491" in panel.html and "202" in panel.html


def test_delimited_text_becomes_a_table():
    """`looks_delimited` is the extraction cascade's own test, reused so the viewer and the parser
    agree about what counts as a table."""
    panel = attachment_view.render(b"PO,Spec,Qty\n208491,STE-402,11\n", "tracker.csv", "")
    assert "<table" in panel.html
    assert panel.label == "Delimited text"


def test_plain_text_is_shown_as_text():
    panel = attachment_view.render(b"just a note about the delivery", "note.txt", "")
    assert "att-text" in panel.html


def test_long_text_is_clipped_and_says_so():
    panel = attachment_view.render(b"x" * (attachment_view.MAX_TEXT_CHARS + 500), "big.txt", "")
    assert "Showing the first" in panel.html


def test_an_rtf_attachment_is_readable():
    rtf = br"{\rtf1\ansi Received 11 EA against PO 208491.}"
    panel = attachment_view.render(rtf, "notice.rtf", "")
    assert "208491" in panel.html


def test_an_image_streams_rather_than_embedding():
    """Inlining these produced a 20 MB response for one photo-heavy message."""
    panel = attachment_view.render(KIND_FIXTURES[sniff.KIND_IMAGE][1](), "p.jpeg", "",
                                   attachment_url="/ui/mail/attachment?id=x&n=0")
    assert "<img" in panel.html
    assert "data:" not in panel.html


# --- Containers ---------------------------------------------------------------

def test_a_zip_lists_what_is_inside_and_makes_each_openable():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", "PO 208491")
        archive.writestr("pod.txt", "signed")
    panel = attachment_view.render(buffer.getvalue(), "bundle.zip", "",
                                   child_url="/ui/mail/attachment/view?id=x&n=0&child=")
    assert "notes.txt" in panel.html and "pod.txt" in panel.html
    assert "data-frag=" in panel.html, "members are not openable"
    assert len(panel.children) == 2


def test_a_tar_is_listed_even_though_the_pipeline_does_not_expand_one():
    import tarfile

    buffer = io.BytesIO()
    with tarfile.open(fileobj=buffer, mode="w:gz") as archive:
        data = b"PO 208491 received"
        info = tarfile.TarInfo("pod.txt")
        info.size = len(data)
        archive.addfile(info, io.BytesIO(data))
    panel = attachment_view.render(buffer.getvalue(), "bundle.tar.gz", "")
    assert "pod.txt" in panel.html


def test_walking_into_a_container_returns_the_member():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", "PO 208491 received")
    found = attachment_view.child_at(buffer.getvalue(), "bundle.zip", "", [0])
    assert found is not None
    content, name, _mime = found
    assert b"208491" in content and name == "notes.txt"


def test_walking_past_the_end_is_not_an_error():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", "hi")
    assert attachment_view.child_at(buffer.getvalue(), "b.zip", "", [7]) is None
    assert attachment_view.child_at(buffer.getvalue(), "b.zip", "", [-1]) is None


def test_nesting_is_capped():
    """`MAX_CONTAINER_DEPTH` is the pipeline's ceiling and the viewer honours it, so a message
    nested inside itself cannot be walked for ever."""
    from config import settings

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", "hi")
    deep = [0] * (settings.MAX_CONTAINER_DEPTH + 1)
    assert attachment_view.child_at(buffer.getvalue(), "b.zip", "", deep) is None


def test_a_zip_bomb_is_refused_by_the_viewer_too():
    """The viewer must not become the one path that will happily inflate a 200:1 archive —
    it expands through the same adapter and the same budget as ingest."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("bomb.txt", b"0" * (60 * 1024 * 1024))
    panel = attachment_view.render(buffer.getvalue(), "bomb.zip", "")
    assert "oversize" in panel.html.lower()


def test_an_unreadable_member_is_listed_with_its_reason_not_hidden():
    """Silently omitting it would be the container equivalent of the blank panel this file
    exists to prevent."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("ok.txt", "fine")
        archive.writestr("huge.bin", b"0" * (60 * 1024 * 1024))
    panel = attachment_view.render(buffer.getvalue(), "bundle.zip", "")
    assert "ok.txt" in panel.html and "huge.bin" in panel.html


# --- Nested mail --------------------------------------------------------------

def eml_bytes(subject="Delivered Notification", body="Received 202 YD against PO 210634",
              attachment=None) -> bytes:
    from email.message import EmailMessage

    message = EmailMessage()
    message["From"] = "warehouse@vendor.example"
    message["To"] = "receiver@premierpm.com"
    message["Subject"] = subject
    message["Date"] = "Tue, 12 Aug 2026 09:00:00 +0000"
    message.set_content(body)
    if attachment is not None:
        # Bytes with an explicit maintype: `add_attachment` picks a handler from the object's
        # type, and the text handler does not take `maintype`.
        message.add_attachment(attachment.encode("utf-8"), maintype="text", subtype="plain",
                               filename="pod.txt")
    return message.as_bytes()


def test_an_attached_email_renders_as_a_message():
    """This is where Premier's delivery evidence actually lives — the POD date, the signature, the
    carrier and the real received quantity all arrived inside a forwarded notification."""
    panel = attachment_view.render(eml_bytes(), "forwarded.eml", "message/rfc822")
    assert panel.label == "Email message"
    assert "Delivered Notification" in panel.html
    assert "warehouse@vendor.example" in panel.html
    assert "210634" in panel.html


def test_an_attached_email_lists_its_own_attachments():
    panel = attachment_view.render(eml_bytes(attachment="signed by Miguel C."), "f.eml",
                                   "message/rfc822",
                                   child_url="/ui/mail/attachment/view?id=x&n=0&child=")
    assert "pod.txt" in panel.html
    assert any(c.filename == "pod.txt" for c in panel.children)


def test_an_attached_email_is_read_as_mime_regardless_of_its_extension():
    """Graph returns an `itemAttachment` as RFC-822 MIME even when the filename says `.msg`. That
    mismatch once emptied `pod_stated_date`, `carrier_name`, `tracking_number` and `received_by`
    in thirteen of thirteen records, so the format is decided by the bytes, never the extension."""
    panel = attachment_view.render(eml_bytes(), "forwarded.msg", "")
    assert panel.label == "Email message"
    assert "Delivered Notification" in panel.html


def test_a_nested_message_body_cannot_run_scripts():
    panel = attachment_view.render(eml_bytes(body="see attached"), "f.eml", "message/rfc822")
    if "sandbox=" in panel.html:
        assert "allow-scripts" not in panel.html


# --- Helpers ------------------------------------------------------------------

def _fixture(kind: str):
    filename, make = KIND_FIXTURES[kind]
    return make(), filename, ""


def test_empty_bytes_do_not_pretend_to_be_a_file():
    panel = attachment_view.render(b"", "gone.pdf", "")
    assert "no bytes" in panel.html
