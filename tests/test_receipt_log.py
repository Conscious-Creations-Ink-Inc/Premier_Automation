"""The receiver report: grouping, and the columns that must stay empty.

This sheet is read as authoritative — it is laid out to sit beside Spitfire's own Receipt Log. So
the two things worth pinning are that rows group into the right lines, and that the three columns
we cannot know are never quietly filled with a plausible number.
"""

from io import BytesIO

import openpyxl

from pipeline import extracted_records_store, receipt_log, state_db
from pipeline.models import ExtractedRecord

NOW = "2026-08-03T12:00:00Z"


def new_conn():
    return state_db.get_connection(":memory:")


def record(**overrides) -> ExtractedRecord:
    values = dict(
        source_email_id="msg-1", po_number="908491", shipment_number=None,
        spec_code="STE-402-LT-B", parent_spec_code="STE-402-LT", sub_spec_suffix="B",
        item_description="BASE, Floor Lamp 2", vendor_name="Light Annex",
        carrier_name="Custom Companies", tracking_number="99942177", quantity_received=1.0,
        unit_of_measure="EA", pod_stated_date="2025-10-09",
        email_date="2025-10-09T23:17:49+05:30", delivery_location=None, comments=None,
        extraction_source="authority_inbound", extraction_confidence=1.0, raw_snippet="…",
        po_line_number=300, received_by="Jordan T.", notification_number="939475",
    )
    values.update(overrides)
    return ExtractedRecord(**values)


def test_build_groups_rows_into_po_then_line_then_receipts():
    conn = new_conn()
    extracted_records_store.write_pending(conn, record(quantity_received=1.0), NOW)
    extracted_records_store.write_pending(conn, record(quantity_received=2.0), NOW)

    report = receipt_log.build(conn)
    [po] = report.purchase_orders
    assert po.po_number == "908491"
    assert po.vendor == "Light Annex"
    [line] = po.lines
    assert line.label == "0300"
    assert line.received == 3.0, "two receipts against one line sum onto it"
    assert len(line.receipts) == 2


def test_two_line_numbers_are_two_lines():
    """Sub-parts are separate Spitfire lines — STE-402-LT-B and -SH are 300 and 301. Merging them
    on their shared parent would report one line with a quantity neither of them has."""
    conn = new_conn()
    extracted_records_store.write_pending(conn, record(po_line_number=300, spec_code="STE-402-LT-B"), NOW)
    extracted_records_store.write_pending(conn, record(po_line_number=301, spec_code="STE-402-LT-SH"), NOW)

    [po] = receipt_log.build(conn).purchase_orders
    assert [line.label for line in po.lines] == ["0300", "0301"]


def test_the_receipt_reference_falls_back_through_three_sources():
    conn = new_conn()
    extracted_records_store.write_pending(conn, record(notification_number="939475"), NOW)
    extracted_records_store.write_pending(
        conn, record(po_line_number=301, notification_number=None, shipment_number="90052"), NOW)
    extracted_records_store.write_pending(
        conn, record(po_line_number=302, notification_number=None, shipment_number=None), NOW)

    refs = [line.receipts[0].reference for line in receipt_log.build(conn).purchase_orders[0].lines]
    assert refs == ["939475", "90052", "Email confirmation"]


def test_order_qty_and_final_are_never_populated_from_email():
    """Order Qty lives on the purchase order inside Spitfire and no delivery email carries it. A
    sheet in this layout reads as authoritative, so an invented order quantity is worse than a gap.

    `Final` is separate: it is a flag someone sets in Spitfire, not a state we could infer.
    """
    conn = new_conn()
    extracted_records_store.write_pending(conn, record(), NOW)
    [line] = receipt_log.build(conn).purchase_orders[0].lines
    assert line.order_qty is None
    assert line.final is False
    assert line.net is None, "Net cannot be known while Order Qty is not"


def test_net_is_derived_from_order_qty_rather_than_stored():
    """`Net = Order Qty - Received` held for all 486 line rows in Premier's own export, so it is
    not an independent fact and must not need a source of its own. Computing it here is what makes
    the column fill itself the day the Spitfire read lands."""
    assert receipt_log.Line(1, "SP", "d", received=2, order_qty=5).net == 3
    assert receipt_log.Line(1, "SP", "d", received=5, order_qty=5).net == 0
    # Over-delivery is real — the corpus has 202 received against 196 ordered — and must show as a
    # negative rather than being clamped, or the sheet hides the discrepancy it exists to surface.
    assert receipt_log.Line(1, "SP", "d", received=7, order_qty=5).net == -2
    # A line with no receipts yet is outstanding in full, not zero.
    assert receipt_log.Line(1, "SP", "d", received=None, order_qty=5).net == 5


def test_final_is_not_inferred_from_quantities():
    """The regression guard. In Premier's export 118 fully-received lines carry no `*` and 5
    partially-received lines do, so any received-versus-ordered rule would be wrong 123 times in
    486. If someone later "fixes" Final by computing it, this fails."""
    fully_received = receipt_log.Line(1, "SP", "d", received=5, order_qty=5)
    assert fully_received.final is False


def test_to_html_returns_a_plain_string_both_uis_can_wrap():
    """Not either UI's `Raw`: the console and /ui each have their own marker class, and this module
    must depend on neither. Escaping still happens here."""
    conn = new_conn()
    extracted_records_store.write_pending(conn, record(item_description='Lamp <b>"A&B"</b>'), NOW)
    out = receipt_log.to_html(receipt_log.build(conn))

    assert type(out) is str
    # `to_html` emits no <b> of its own, so any that survived came from the description.
    assert "<b>" not in out
    assert "&lt;b&gt;" in out and "&amp;" in out and "&quot;" in out


def test_every_preview_row_is_exactly_as_wide_as_the_header():
    """A row one cell short does not fail loudly — every column right of the gap shifts left by
    one and the table still renders, which is exactly how a receipt date came to sit under a
    heading that said Order Qty. Counting `colspan` here is what makes that a test failure rather
    than something noticed in a screenshot."""
    import re

    conn = new_conn()
    extracted_records_store.write_pending(conn, record(), NOW)
    out = receipt_log.to_html(receipt_log.build(conn))

    rows = re.findall(r"<tr[^>]*>(.*?)</tr>", out)
    assert rows, "the preview rendered no rows at all"
    for row in rows:
        width = sum(int(m or 1) for m in re.findall(r'<td(?:[^>]*?colspan="(\d+)")?[^>]*>', row))
        assert width == receipt_log._HTML_COLUMNS, f"row is {width} cells wide: {row}"


def test_the_preview_never_puts_a_date_in_the_order_qty_column():
    """Spitfire's own sheet reuses Order Qty (col F) for a receipt's date. The preview splits Date
    out instead, so Order Qty holds order quantities and nothing else. `to_xlsx` still follows
    Spitfire — see `test_receipt_log_conformance.py`."""
    import re

    conn = new_conn()
    extracted_records_store.write_pending(conn, record(), NOW)
    out = receipt_log.to_html(receipt_log.build(conn))

    # `[^>]*` because a row also carries `data-group`, which is what lets the UI page this sheet by
    # purchase order rather than splitting one across a page boundary. Matching the tag exactly made
    # this test fail on a markup addition it has no opinion about.
    header = re.search(r'<tr class="r-head"[^>]*>(.*?)</tr>', out).group(1)
    columns = re.findall(r"<td[^>]*>(.*?)</td>", header)
    assert columns.index("Date") < columns.index("Order Qty")

    receipt = re.search(r'<tr class="r-rcpt"[^>]*>(.*?)</tr>', out).group(1)
    cells = re.findall(r"<td[^>]*>(.*?)</td>", receipt)
    assert re.match(r"\d{4}-\d{2}-\d{2}", cells[columns.index("Date")])
    assert cells[columns.index("Order Qty")] == ""


def test_the_sheet_title_matches_premiers_own_file():
    """Premier's `_Spitfire-LIVE_General_Receipt Log.xlsx` reads 'Premier Design to Completion
    Report' in H1. The word 'Report' was missing here."""
    assert receipt_log.SHEET_TITLE == "Premier Design to Completion Report"


def test_to_xlsx_produces_a_workbook_with_the_headers_on_row_five():
    conn = new_conn()
    extracted_records_store.write_pending(conn, record(), NOW)
    workbook = openpyxl.load_workbook(BytesIO(receipt_log.to_xlsx(receipt_log.build(conn))))

    sheet = workbook["Receipt Log"]
    assert sheet["H1"].value == receipt_log.SHEET_TITLE
    assert sheet["H2"].value == receipt_log.SHEET_SUBTITLE
    assert [sheet.cell(5, col).value for col in (1, 2, 3, 4)] == ["DocNo", "Vendor", "Line", "Description"]
    # The three unknowable columns carry a header and never a value.
    assert sheet.cell(5, 6).value == "Order Qty"
    assert all(sheet.cell(row, 6).value in (None, "On") or hasattr(sheet.cell(row, 6).value, "year")
               for row in range(7, sheet.max_row + 1))


def test_an_empty_store_still_renders_both_ways():
    """A fresh checkout must produce a sheet that says so rather than raising."""
    conn = new_conn()
    report = receipt_log.build(conn)
    assert report.is_empty
    assert report.empty_message in receipt_log.to_html(report)
    assert receipt_log.to_xlsx(report)[:2] == b"PK"


def test_build_does_not_require_the_caller_to_set_a_row_factory():
    """`api.deps.get_pipeline_conn` does not set one. Without `build` setting it itself the failure
    is a TypeError deep in the grouping loop that names nothing useful."""
    conn = new_conn()
    conn.row_factory = None
    extracted_records_store.write_pending(conn, record(), NOW)
    assert receipt_log.build(conn).purchase_orders[0].po_number == "908491"
    assert conn.row_factory is None, "the caller's factory is restored"


# --- what the purchase order fills in -----------------------------------------------------------
#
# Three of this sheet's columns describe the order, not the delivery, and no email carries them.
# They were blank on every row until `spitfire_po_lines` could answer them.


def _store_with_a_record(**overrides):
    """One record and one mirrored PO line, in a fresh in-memory store."""
    from pipeline import state_db

    conn = state_db.get_connection(":memory:")
    fields = dict(source_email_id="mail-1", po_number="908491", spec_code="STE-402-LT-B",
                  item_description="BASE, Floor Lamp 2", vendor_name=None,
                  quantity_received=11.0, unit_of_measure=None, pod_stated_date="2025-10-01",
                  received_by=None, po_line_number=300, origin="auto", created_by=None,
                  email_date="2025-10-01", extraction_source="test", extraction_confidence=1.0)
    fields.update(overrides)
    columns = ", ".join(fields)
    conn.execute(
        f"INSERT INTO extracted_records ({columns}, created_at) "
        f"VALUES ({', '.join('?' * len(fields))}, 'now')", tuple(fields.values()))
    conn.execute(
        """INSERT INTO spitfire_po_lines
           (line_key, po_number, line_number, spec_code, description, vendor_name,
            unit_of_measure, qty_ordered, qty_received, qty_in_transit, refreshed_at)
           VALUES ('k1', '908491', 300, 'STE-402-LT-B', 'BASE, Floor Lamp 2', 'Light Annex',
                   'EA', 12.0, 0.0, 0.0, 'now')""")
    conn.commit()
    return conn


def test_order_qty_is_filled_from_the_mirrored_purchase_order():
    """And Net lights up on its own, because it is `Order Qty - Received` and nothing else."""
    report = receipt_log.build(_store_with_a_record())
    line = report.purchase_orders[0].lines[0]

    assert line.order_qty == 12.0
    assert line.received == 11.0
    assert line.net == 1.0


def test_vendor_and_uom_are_filled_from_the_purchase_order_too():
    """`completeness.DERIVED` says a person must never be asked to type these. This is where they
    come from instead."""
    report = receipt_log.build(_store_with_a_record())

    assert report.purchase_orders[0].vendor == "Light Annex"
    assert report.purchase_orders[0].lines[0].uom == "EA"


def test_what_the_record_already_says_is_never_overwritten():
    """Fallback only. A vendor or a unit the record carries is what a reviewer has been looking at,
    and silently replacing it would change what they thought they were approving."""
    report = receipt_log.build(
        _store_with_a_record(vendor_name="P. Kaufmann", unit_of_measure="YD"))

    assert report.purchase_orders[0].vendor == "P. Kaufmann"
    assert report.purchase_orders[0].lines[0].uom == "YD"


def test_a_store_with_no_mirrored_lines_builds_exactly_as_before():
    """The safety property. Nothing about this join may change a report on a store that cannot
    answer it — which is every store until a purchase order has been pulled."""
    conn = _store_with_a_record()
    conn.execute("DELETE FROM spitfire_po_lines")
    conn.commit()

    line = receipt_log.build(conn).purchase_orders[0].lines[0]
    assert line.order_qty is None and line.net is None


def test_a_line_matched_by_spec_when_no_line_number_was_stated():
    """Most mail states no line number, so matching on the line alone would leave the columns blank
    for nearly all of it."""
    line = receipt_log.build(_store_with_a_record(po_line_number=None)).purchase_orders[0].lines[0]
    assert line.order_qty == 12.0


def test_final_is_never_filled_by_the_join():
    """It is a flag a person sets in Spitfire. 118 fully received lines in Premier's own export
    carry no asterisk and 5 partially received ones do, so no rule over quantities can produce it —
    and this line is fully received, which is exactly where the temptation lies."""
    report = receipt_log.build(_store_with_a_record(quantity_received=12.0))
    assert report.purchase_orders[0].lines[0].final is False


# --- the manual / automated flag ----------------------------------------------------------------
#
# It rides the Receiver column, which already held a word rather than a name. That is what keeps
# the sheet at its exact nine columns and conformant with Premier's template.


def test_an_automated_record_reads_as_automation():
    report = receipt_log.build(_store_with_a_record())
    assert report.purchase_orders[0].lines[0].receipts[0].receiver == "Automation"


def test_a_manual_record_names_the_person_who_entered_it():
    report = receipt_log.build(
        _store_with_a_record(origin="manual", created_by="M Rivera"))
    assert report.purchase_orders[0].lines[0].receipts[0].receiver == "M Rivera"


def test_whoever_signed_for_the_goods_outranks_both():
    """A named signature is the more specific truth about who took delivery, whichever way the
    record was made."""
    report = receipt_log.build(
        _store_with_a_record(origin="manual", created_by="M Rivera", received_by="U ALI"))
    assert report.purchase_orders[0].lines[0].receipts[0].receiver == "U ALI"


def test_the_flag_adds_no_column_to_the_sheet():
    """The whole reason it rides Receiver. A tenth column would take the workbook out of
    conformance with Premier's own export, which `test_receipt_log_conformance` holds cell for
    cell."""
    assert len(receipt_log._HEADERS) == 9
    assert [name.strip() for _, name in receipt_log._HEADERS] == [
        "DocNo", "Vendor", "Line", "Description", "Order Qty", "Received", "Net", "Final",
        "Receiver"]


def test_manual_and_automated_records_produce_the_same_report_shape():
    """Same columns, same layout, same row types — differing only in the one cell that says which.
    A person reading the PDF can tell them apart; a machine diffing the structure cannot."""
    from io import BytesIO

    import openpyxl

    def sheet(**overrides):
        book = openpyxl.load_workbook(
            BytesIO(receipt_log.to_xlsx(receipt_log.build(_store_with_a_record(**overrides)))))
        return book["Receipt Log"]

    automated = sheet()
    manual = sheet(origin="manual", created_by="M Rivera")

    assert automated.max_row == manual.max_row
    assert automated.max_column == manual.max_column
    differing = [(r, c) for r in range(1, automated.max_row + 1)
                 for c in range(1, automated.max_column + 1)
                 if automated.cell(r, c).value != manual.cell(r, c).value]
    assert len(differing) == 1, differing
    row, column = differing[0]
    assert automated.cell(row, column).value == "Automation"
    assert manual.cell(row, column).value == "M Rivera"
