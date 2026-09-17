"""Every attachment kind reaches an adapter and ends with a recorded verdict.

The promise is *"every file arriving as an Outlook attachment is analysed, and nothing is ever
silently dropped"*. That is only worth anything if it stays true as formats are added, so it is
enforced structurally rather than by inspection:

* `test_fixture_table_covers_every_sniffable_kind` fails the moment a kind is added to
  `sniff.ALL_KINDS` without a fixture here — you cannot introduce a format without deciding
  what reads it.
* `test_every_kind_reaches_a_terminal_disposition` proves each one ends somewhere real. Note
  that "analysed" includes an honest `unsupported_format`: a `.doc` gets a type, a ledger row
  and a sentence telling the operator what to do, which is the correct answer when no reliable
  reader exists. What is *not* acceptable is `observed` — an attachment that vanished.
"""

import io
import struct
import zipfile

import pytest

from pipeline import attachment_ledger, ingest_orchestrator, state_db
from pipeline.models import Attachment, RawEmail
from pipeline.parsing import sniff
from pipeline.stage3_extract import containers, dispatch
from pipeline.stage3_extract.base import ExtractionSource

NOW = "2026-08-03T00:00:00Z"


# --- Fixture bytes, one per sniffable kind ----------------------------------


def make_pdf() -> bytes:
    from tests.test_extract import make_pdf_bytes
    return make_pdf_bytes([["PO", "Spec", "Qty"], ["908491", "STE-402-LT-B", "11"]])


def make_xlsx() -> bytes:
    import openpyxl
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Vendor", "PO#", "Spec#", "QTY", "Item Description", "Confirmed Received: Yes or No"])
    sheet.append(["Daniel Stuart", "907030", "LOB-203-PI", 12, '18"x18" Throw Pillow', "yes"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_xls() -> bytes:
    """A minimal BIFF workbook. xlwt is not a dependency, so this is hand-built: an OLE2
    container whose directory names the `Workbook` stream is enough for the sniffer, and the
    adapter's failure to parse it is itself a valid, recorded outcome."""
    ole_header = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 24
    workbook_stream_name = "Workbook".encode("utf-16-le")
    return ole_header + b"\x00" * 100 + workbook_stream_name + b"\x00" * 400


def make_doc() -> bytes:
    ole_header = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 24
    return ole_header + b"\x00" * 100 + "WordDocument".encode("utf-16-le") + b"\x00" * 400


def make_docx() -> bytes:
    import docx
    document = docx.Document()
    document.add_paragraph("Received 11 EA of STE-402-LT-B against PO 908491.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pptx() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("ppt/presentation.xml", "<presentation/>")
    return buffer.getvalue()


def make_zip() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", "Received against PO 908491, spec STE-402-LT-B, 11 EA.")
    return buffer.getvalue()


def make_zip_office_plain() -> bytes:
    """A zip that is neither an OOXML document nor something we expand into records."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("readme.md", "nothing useful here")
    return buffer.getvalue()


def make_msg() -> bytes:
    ole_header = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 24
    return ole_header + b"\x00" * 64 + b"__substg1.0_0037001F" + b"\x00" * 400


def make_image() -> bytes:
    """A JPEG large enough to be treated as evidence rather than chrome."""
    return b"\xff\xd8\xff\xe0" + b"\x00" * (200 * 1024)


def make_html() -> bytes:
    return (b"<html><body><table>"
            b"<tr><th>PO</th><th>Spec</th><th>Qty</th></tr>"
            b"<tr><td>908491</td><td>STE-402-LT-B</td><td>11</td></tr>"
            b"</table></body></html>")


def make_text() -> bytes:
    return b"Vendor,PO#,Spec#,QTY\nDaniel Stuart,907030,LOB-203-PI,12\n"


def make_archive() -> bytes:
    import gzip
    return gzip.compress(b"Received against PO 908491")


def make_archive_unsupported() -> bytes:
    return b"7z\xbc\xaf\x27\x1c" + b"\x00" * 200


def make_legacy_office() -> bytes:
    """OLE2 that names neither Word nor Excel nor Outlook."""
    return b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 500


def make_unknown() -> bytes:
    return b"\x00\x01\x02\x03\xfe\xff" + bytes(range(200, 256)) * 4


KIND_FIXTURES = {
    sniff.KIND_PDF: ("pod.pdf", make_pdf),
    sniff.KIND_XLSX: ("tracker.xlsx", make_xlsx),
    sniff.KIND_XLS: ("tracker.xls", make_xls),
    sniff.KIND_DOC: ("notice.doc", make_doc),
    sniff.KIND_DOCX: ("notice.docx", make_docx),
    sniff.KIND_PPTX: ("deck.pptx", make_pptx),
    sniff.KIND_ZIP_OFFICE: ("bundle.zip", make_zip_office_plain),
    sniff.KIND_MSG: ("forwarded.msg", make_msg),
    sniff.KIND_IMAGE: ("IMG_2479.jpeg", make_image),
    sniff.KIND_HTML: ("notice.html", make_html),
    sniff.KIND_TEXT: ("tracker.csv", make_text),
    sniff.KIND_ARCHIVE: ("bundle.tar.gz", make_archive),
    sniff.KIND_ARCHIVE_UNSUPPORTED: ("bundle.7z", make_archive_unsupported),
    sniff.KIND_LEGACY_OFFICE: ("mystery.bin", make_legacy_office),
    sniff.KIND_UNKNOWN: ("mystery.dat", make_unknown),
}


def test_fixture_table_covers_every_sniffable_kind():
    """Adding a kind to the sniffer without deciding what reads it fails here.

    This is the mechanism that keeps "every attachment is analysed" honest over time, rather
    than true only on the day it was written.
    """
    assert set(KIND_FIXTURES) == set(sniff.ALL_KINDS), (
        f"kinds with no fixture: {sorted(set(sniff.ALL_KINDS) - set(KIND_FIXTURES))}; "
        f"fixtures for unknown kinds: {sorted(set(KIND_FIXTURES) - set(sniff.ALL_KINDS))}"
    )


@pytest.fixture
def conn():
    connection = state_db.get_connection(":memory:")
    yield connection
    connection.close()


@pytest.mark.parametrize("kind", sorted(KIND_FIXTURES))
def test_every_kind_reaches_a_terminal_disposition(kind, conn):
    filename, factory = KIND_FIXTURES[kind]
    data = factory()

    assert sniff.sniff(data, filename).kind == kind, "fixture does not sniff as the kind it claims"

    email = RawEmail(
        email_id=f"msg-{kind}", received_at=NOW, sender_address="a@b.com", sender_domain="b.com",
        subject="Delivery confirmation", body_html=None, body_text=None,
        attachments=[Attachment(filename=filename, content_type="", content_bytes=data,
                                sniffed_kind=kind, size_bytes=len(data))],
    )
    attachment_ledger.observe(conn, email, "hold", NOW)
    attachment = email.attachments[0]
    assert attachment.ledger_id is not None, "every attachment must be ledgered at ingest"

    dispatch.dispatch_source(
        conn,
        ExtractionSource(
            source_email_id=email.email_id, email_date=NOW, source_type="attachment",
            filename=filename, content_type="", content_bytes=data,
            ledger_id=attachment.ledger_id,
        ),
        ingest_orchestrator.build_default_adapters(),
        budget=containers.Budget.fresh(),
        now=NOW,
    )

    rows = attachment_ledger.for_email(conn, email.email_id)
    top = next(r for r in rows if r.id == attachment.ledger_id)
    assert top.disposition in attachment_ledger.ALL_DISPOSITIONS
    assert top.disposition in attachment_ledger.TERMINAL, (
        f"{kind} was left at {top.disposition!r} — it escaped the dispatcher without a verdict"
    )
    assert top.sniffed_kind == kind
    assert attachment_ledger.orphans(conn) == []


@pytest.mark.parametrize("kind", sorted(KIND_FIXTURES))
def test_no_kind_is_left_unclaimed(kind, conn):
    """`no_adapter` means a genuine coverage gap — distinct from `unsupported_format`, which is
    a decision we made and can explain."""
    filename, factory = KIND_FIXTURES[kind]
    data = factory()
    source = ExtractionSource(
        source_email_id="m", email_date=NOW, source_type="attachment",
        filename=filename, content_type="", content_bytes=data,
    )
    adapters = ingest_orchestrator.build_default_adapters()
    claimed = [a.__class__.__name__ for a in adapters if _safely_claims(a, source)]
    expandable = [c.__class__.__name__ for c in containers.build_default_containers()
                  if c.can_expand(source)]
    assert claimed or expandable, f"nothing claims {kind}"


def _safely_claims(adapter, source) -> bool:
    try:
        return adapter.can_handle(source)
    except Exception:
        return False


def test_unsupported_formats_carry_operator_guidance(conn):
    """A `.doc` is answered, not merely refused — the ledger holds a sentence a person can act
    on rather than a silent absence."""
    data = make_doc()
    email = RawEmail(
        email_id="msg-doc", received_at=NOW, sender_address="a@b.com", sender_domain="b.com",
        subject="Delivery", body_html=None, body_text=None,
        attachments=[Attachment(filename="notice.doc", content_type="", content_bytes=data,
                                sniffed_kind=sniff.KIND_DOC, size_bytes=len(data))],
    )
    attachment_ledger.observe(conn, email, "hold", NOW)
    dispatch.dispatch_source(
        conn,
        ExtractionSource(source_email_id="msg-doc", email_date=NOW, source_type="attachment",
                         filename="notice.doc", content_type="", content_bytes=data,
                         ledger_id=email.attachments[0].ledger_id),
        ingest_orchestrator.build_default_adapters(), budget=containers.Budget.fresh(), now=NOW,
    )
    row = attachment_ledger.for_email(conn, "msg-doc")[0]
    assert row.disposition == attachment_ledger.UNSUPPORTED_FORMAT
    assert ".docx" in row.disposition_detail
    assert row.review_status == attachment_ledger.REVIEW_PENDING


def test_zip_members_are_expanded_and_individually_ledgered(conn):
    data = make_zip()
    email = RawEmail(
        email_id="msg-zip", received_at=NOW, sender_address="a@b.com", sender_domain="b.com",
        subject="Delivery", body_html=None, body_text=None,
        attachments=[Attachment(filename="bundle.zip", content_type="", content_bytes=data,
                                sniffed_kind=sniff.KIND_ZIP_OFFICE, size_bytes=len(data))],
    )
    attachment_ledger.observe(conn, email, "hold", NOW)
    dispatch.dispatch_source(
        conn,
        ExtractionSource(source_email_id="msg-zip", email_date=NOW, source_type="attachment",
                         filename="bundle.zip", content_type="", content_bytes=data,
                         ledger_id=email.attachments[0].ledger_id),
        ingest_orchestrator.build_default_adapters(), budget=containers.Budget.fresh(), now=NOW,
    )
    rows = attachment_ledger.for_email(conn, "msg-zip")
    assert len(rows) == 2, "the archive and its member each need their own row"
    container_row = next(r for r in rows if r.depth == 0)
    member_row = next(r for r in rows if r.depth == 1)
    assert container_row.disposition == attachment_ledger.CONTAINER_EXPANDED
    assert member_row.filename == "notes.txt"
    assert member_row.container_path == "bundle.zip!/notes.txt"
    assert attachment_ledger.orphans(conn) == []


def test_a_zip_bomb_is_refused_not_expanded(conn):
    """Ten megabytes of zeroes compress to a few kilobytes. The ratio guard reads the archive's
    own metadata, so nothing is decompressed to discover it was too big."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("bomb.bin", b"\0" * 10_000_000)
    data = buffer.getvalue()
    assert len(data) < 100_000, "fixture should be highly compressed"

    email = RawEmail(
        email_id="msg-bomb", received_at=NOW, sender_address="a@b.com", sender_domain="b.com",
        subject="Delivery", body_html=None, body_text=None,
        attachments=[Attachment(filename="bomb.zip", content_type="", content_bytes=data,
                                sniffed_kind=sniff.KIND_ZIP_OFFICE, size_bytes=len(data))],
    )
    attachment_ledger.observe(conn, email, "hold", NOW)
    records = dispatch.dispatch_source(
        conn,
        ExtractionSource(source_email_id="msg-bomb", email_date=NOW, source_type="attachment",
                         filename="bomb.zip", content_type="", content_bytes=data,
                         ledger_id=email.attachments[0].ledger_id),
        ingest_orchestrator.build_default_adapters(), budget=containers.Budget.fresh(), now=NOW,
    )
    assert records == []
    member = next(r for r in attachment_ledger.for_email(conn, "msg-bomb") if r.depth == 1)
    assert member.disposition == attachment_ledger.DROPPED_OVERSIZE
    assert "compression ratio" in member.disposition_detail


def test_a_forwarded_notification_expands_whether_it_is_named_eml_or_msg():
    """The same bytes arrive under both extensions and must read identically.

    Graph hands a forwarded message over as MIME with `contentType: message/rfc822`; whether the
    connector names it `.eml` or `.msg` is incidental. Premier's mailbox holds one message where
    the `.msg` copy expanded into its two MIME parts and the `.eml` copy — byte-for-byte identical,
    same sha256 — was handed to the OLE reader and died `BadZipFile: File is not a zip file`,
    taking a Delivered Notification with it. Deciding by extension is what makes that possible, so
    this asserts the extension cannot decide anything.
    """
    from pipeline.stage3_extract.base import ExtractionSource
    from pipeline.stage3_extract.containers import Budget, MsgContainerAdapter

    raw = (
        b"Received: from LV8PR14MB7645.namprd14.prod.outlook.com (2603::1) by mx.example\r\n"
        b"From: routing@example-logistics.test\r\n"
        b"Subject: [External] 99985 - Delivered Notification - 910634\r\n"
        b"MIME-Version: 1.0\r\n"
        b'Content-Type: multipart/alternative; boundary="b1"\r\n'
        b"\r\n--b1\r\nContent-Type: text/plain; charset=utf-8\r\n\r\n"
        b"Authority #: 99985\r\nDelivered: 09/10/2025 Signed by: U ALI\r\n"
        b"\r\n--b1\r\nContent-Type: text/html; charset=utf-8\r\n\r\n"
        b"<html><body><p>Authority #: 99985</p></body></html>\r\n--b1--\r\n"
    )

    expanded = {}
    for name in ("notification.eml", "notification.msg"):
        source = ExtractionSource(
            source_email_id="msg-1", email_date="2026-08-12", source_type="attachment",
            filename=name, content_bytes=raw, content_type="message/rfc822",
        )
        adapter = MsgContainerAdapter()
        assert adapter.can_expand(source), name
        expanded[name] = [child.filename for child in adapter.expand(source, Budget.fresh())]

    assert expanded["notification.eml"], "MIME bytes named .eml must still expand"
    assert expanded["notification.eml"] == expanded["notification.msg"]
