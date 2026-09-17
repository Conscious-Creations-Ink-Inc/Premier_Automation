import io
import zipfile
from typing import List, Optional

import docx

from pipeline.models import ExtractedRecord
from pipeline.parsing import sniff
from pipeline.stage3_extract.base import (
    ExtractionAdapter,
    ExtractionSource,
    build_record_from_row,
    is_item_row,
    map_headers,
)
from pipeline.stage3_extract.freetext_adapter import FreetextAdapter
from pipeline.stage3_extract.ocr_adapter import (
    DocumentIntelligenceClient,
    MockDocumentIntelligenceClient,
    analyze_with_retry,
    records_from_ocr_result,
)

DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _extract_embedded_images(content_bytes: bytes) -> List[bytes]:
    """Word docs store any pasted photo/scan under word/media/ inside the zip package — the
    real-world shape for a photographed POD pasted directly into a Word document."""
    with zipfile.ZipFile(io.BytesIO(content_bytes)) as z:
        return [z.read(name) for name in z.namelist() if name.startswith("word/media/")]


class DocxAdapter(ExtractionAdapter):
    """Word documents — not in the original discovery-doc sample set, but a real, plausible
    delivery-paperwork format found in real test samples. Tries a native table first (same
    fuzzy-header matching as Html/Pdf/Excel); falls back to any embedded image (a photographed
    POD pasted into the doc) via OCR; falls back to the document's own paragraph text via
    FreetextAdapter as a last resort — mirroring PdfAdapter's table -> text-layer -> nothing
    cascade for a document format that can arrive in any of those three shapes."""

    def __init__(self, ocr_client: Optional[DocumentIntelligenceClient] = None):
        self.ocr_client = ocr_client or MockDocumentIntelligenceClient()

    def can_handle(self, source: ExtractionSource) -> bool:
        """Claim by byte sniff, not by declared content type.

        This was the one adapter still comparing `content_type` as a string, so a .docx arriving
        as `application/octet-stream` — or with no type at all, which is how the corpus's real
        attachments arrive — reached no adapter. Same rationale as ExcelAdapter's.
        """
        if source.source_type != "attachment" or source.content_bytes is None:
            return False
        return sniff.sniff(source.content_bytes, source.filename or "",
                           source.content_type or "").kind == sniff.KIND_DOCX

    def extract(self, source: ExtractionSource) -> List[ExtractedRecord]:
        document = docx.Document(io.BytesIO(source.content_bytes))

        records: List[ExtractedRecord] = []
        for table in document.tables:
            cells_grid = [[cell.text for cell in row.cells] for row in table.rows]
            if len(cells_grid) < 2:
                continue
            column_map = map_headers(cells_grid[0])
            if not column_map:
                continue
            records.extend(build_record_from_row(source, row, column_map, "docx")
                           for row in cells_grid[1:] if is_item_row(row))
        if records:
            return records

        for image_bytes in _extract_embedded_images(source.content_bytes):
            # `analyze_with_retry` raises `OcrServiceUnavailable` when the service will not answer,
            # and that is deliberately allowed to propagate: `dispatch` records it against this
            # attachment so the outage is re-runnable. It used to be caught here and turned into a
            # placeholder record, which made the ledger claim a successful extraction.
            result = analyze_with_retry(self.ocr_client, image_bytes, source.source_email_id)
            records.extend(records_from_ocr_result(source, result, extraction_source="docx"))
        if records:
            return records

        text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
        text_source = ExtractionSource(
            source_email_id=source.source_email_id, email_date=source.email_date,
            source_type="body", body_text=text,
            ledger_id=source.ledger_id, known_po_numbers=source.known_po_numbers,
        )
        records = FreetextAdapter().extract(text_source)
        for r in records:
            r.extraction_source = "docx"
        return records
