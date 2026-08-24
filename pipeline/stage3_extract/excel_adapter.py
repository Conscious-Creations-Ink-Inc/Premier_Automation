"""Spreadsheet trackers (corpus class C2) — `.xlsx` and legacy `.xls`.

When a property or vendor stops replying, Premier switches to a spreadsheet: one row per item,
with a confirmation column the recipient fills in. `Cameo Receivers.xlsx` (148 rows) and
`Public Space - Pending Receipt Confirmation Orders.xlsx` (92 rows) are the real ones, and
between them they carry more receivable lines than the rest of the corpus combined.

Dispatch is by byte sniff, not content type: the corpus tracker arrives with `mimetype=None`,
so a content-type equality check never fired on it (finding C13).

`.xls` is handled here too rather than in a separate adapter. It is the same document, the same
columns and the same downstream records — only the container differs — and `xlrd` 2.x reads
exactly and only that legacy format.
"""

import io
from typing import List

import openpyxl

from pipeline.models import ExtractedRecord
from pipeline.parsing import integrity, sniff
from pipeline.stage3_extract import grid_reader
from pipeline.stage3_extract.base import ExtractionAdapter, ExtractionSource

EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
"""Kept for callers building fixtures. Dispatch no longer compares against it — the real
trackers arrive with no content type at all."""

MAX_HEADER_SCAN_ROWS = grid_reader.MAX_HEADER_SCAN_ROWS   # re-exported; tests import it


class ExcelAdapter(ExtractionAdapter):
    def can_handle(self, source: ExtractionSource) -> bool:
        if source.source_type != "attachment" or not source.content_bytes:
            return False
        kind = sniff.sniff(source.content_bytes, source.filename or "", source.content_type or "").kind
        return kind in (sniff.KIND_XLSX, sniff.KIND_XLS)

    def extract(self, source: ExtractionSource) -> List[ExtractedRecord]:
        kind = sniff.sniff(source.content_bytes, source.filename or "", source.content_type or "").kind
        if kind == sniff.KIND_XLS:
            return self._extract_xls(source)
        return self._extract_xlsx(source)

    # --- .xlsx ---------------------------------------------------------------

    def _extract_xlsx(self, source: ExtractionSource) -> List[ExtractedRecord]:
        state, detail = integrity.office_state(source.content_bytes)
        if state == integrity.ENCRYPTED:
            raise PermissionError(f"encrypted workbook: {detail}")

        workbook = openpyxl.load_workbook(io.BytesIO(source.content_bytes), data_only=True, read_only=True)
        records: List[ExtractedRecord] = []
        try:
            for sheet in workbook.worksheets:
                rows = [grid_reader.as_strings(row) for row in sheet.iter_rows(values_only=True)]
                records.extend(grid_reader.records_from_rows(source, rows, f"excel:{sheet.title}"))
        finally:
            workbook.close()
        return records

    # --- .xls ----------------------------------------------------------------

    def _extract_xls(self, source: ExtractionSource) -> List[ExtractedRecord]:
        """Legacy BIFF workbooks, via xlrd.

        Dates arrive as floats in Excel's serial-number epoch rather than as datetimes, so they
        are converted here; left alone, `45108.0` would reach `normalize_date` and be discarded,
        losing the delivery date the tracker exists to record.
        """
        import xlrd

        book = xlrd.open_workbook(file_contents=source.content_bytes)
        records: List[ExtractedRecord] = []
        try:
            for sheet in book.sheets():
                rows = []
                for row_index in range(sheet.nrows):
                    cells = []
                    for column_index in range(sheet.ncols):
                        cell = sheet.cell(row_index, column_index)
                        cells.append(_xls_cell_value(cell, book.datemode))
                    rows.append(cells)
                records.extend(grid_reader.records_from_rows(source, rows, f"xls:{sheet.name}"))
        finally:
            book.release_resources()
        return records


def _xls_cell_value(cell, datemode) -> str:
    """One BIFF cell as the string the grid reader expects."""
    import xlrd

    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            parts = xlrd.xldate_as_tuple(cell.value, datemode)
            return f"{parts[0]:04d}-{parts[1]:02d}-{parts[2]:02d}"
        except Exception:
            return str(cell.value)
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        # Excel stores every number as a float; an integral quantity should read "12", not "12.0".
        return str(int(cell.value)) if float(cell.value).is_integer() else str(cell.value)
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return ""
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return "yes" if cell.value else "no"
    return str(cell.value).strip()
