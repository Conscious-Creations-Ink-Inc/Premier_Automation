"""One attachment, opened full size — whatever kind it is.

The promise this file exists to keep is *"every file arriving as an Outlook attachment can be
looked at in the UI"*. Before it, `mail_view._preview` had branches for image, pdf, xlsx/xls and
docx and returned `""` for the other ten kinds the sniffer recognises, so a `.msg` carrying the POD
evidence, an HTML delivery confirmation and a CSV tracker all rendered as a filename and a Download
link — the reader had to leave the browser to find out what they had been sent.

**Nothing here ever executes an attachment.** Two rules, and they are what make "open any type"
safe rather than reckless:

  * Anything we can read, we read server-side and re-emit as our own escaped markup. A `.csv`
    becomes a table *we* built; the sender's bytes never reach the browser as code.
  * Anything we cannot re-emit — an HTML attachment, a nested message body — goes in an iframe with
    no `allow-scripts`, behind `default-src 'none'`. `tests/test_operations_readonly` regexes every
    `sandbox=` in this package and fails if that ever changes.

The raw-bytes route (`/ui/mail/attachment`) keeps its own, stricter rule: only images and PDFs may
be served inline, everything else downloads. This module did not widen that door — it made the door
unnecessary for reading.

`str` in, `str` out, like `mail_view`: this module knows about neither interface's `Raw` class and
does its own escaping, so both can wrap the result.
"""
import io
import re
from dataclasses import dataclass, field
from html import escape
from typing import Callable, Dict, List, Optional, Sequence

from config import settings
from pipeline.parsing import sniff
from pipeline.stage3_extract import containers, text_adapter
from pipeline.stage3_extract.base import ExtractionSource
from pipeline.stage3_extract.unsupported_adapter import GUIDANCE

MAX_TEXT_CHARS = 200_000
"""How much text a panel will show. Generous — the point is to read the thing — but not unbounded:
a 25 MB log attachment would otherwise become a 25 MB HTML response."""

MAX_TABLE_ROWS = 2000
MAX_TABLE_COLS = 40
"""Per sheet. The old preview stopped at 3 sheets x 60 rows x 15 columns, which silently hid most
of a tracker; the pager makes a bigger table cheap to read, so the clip only has to stop a
pathological workbook from becoming an unusable page."""

MAX_MEMBERS = 500


@dataclass
class Panel:
    """A rendered attachment: the markup, plus what the renderer decided it was looking at."""
    html: str
    kind: str
    label: str
    """Human wording for the type, shown in the header — "Spreadsheet", "Email message"."""
    children: List["ChildRef"] = field(default_factory=list)
    """Members of a container, addressable as `child=` on the viewer route."""


@dataclass(frozen=True)
class ChildRef:
    index: int
    filename: str
    size_bytes: int
    note: str = ""


def render(content: bytes, filename: str = "", content_type: str = "", *,
           attachment_url: str = "/mail/attachment", child_url: str = "") -> Panel:
    """Render one attachment. Never raises — a renderer that fails says so and offers the file.

    The kind is decided by **re-sniffing the bytes**, not by trusting the stored `kind` or the
    sender's declared type. Both are unreliable here: `mail_view._from_graph` hardcodes `kind=""`,
    every one of the 35 cached rows in Premier's live store has `kind = ''`, and the sniffer's own
    docstring notes the declared content type is frequently absent. Sniffing is what the extraction
    cascade does for the same reason.
    """
    if not content:
        return Panel(_note("This attachment has no bytes — see the verdict beside it."),
                     kind=sniff.KIND_UNKNOWN, label="Empty")

    result = sniff.sniff(content, filename or "", content_type or "")
    renderer = _RENDERERS.get(result.kind, _unsupported)
    try:
        return renderer(content, filename, content_type, attachment_url, child_url, result.kind)
    except Exception as exc:                                    # noqa: BLE001 — shown to the user
        # A renderer that raises must not take the popup with it. Saying which reader failed and on
        # what is more use than a blank panel, and the Download link in the header still works.
        return Panel(
            _note(f"This {result.kind} could not be read: {type(exc).__name__}: {exc}")
            + _note("The file itself is intact — use Download to open it outside the browser."),
            kind=result.kind, label=_LABELS.get(result.kind, result.kind),
        )


# --- Readable formats ---------------------------------------------------------

def _image(content, filename, content_type, attachment_url, child_url, kind) -> Panel:
    """Streamed from the raw route rather than embedded as a data URI. One photo-heavy message
    produced a 20 MB response when these were inlined — the note on `_resolve_inline_images` in
    `mail_view` records it."""
    return Panel(f'<img class="view-img" src="{_esc(attachment_url)}" alt="{_esc(filename)}">',
                 kind=kind, label="Image")


def _pdf(content, filename, content_type, attachment_url, child_url, kind) -> Panel:
    """The browser's own PDF viewer, which brings paging, search and zoom for free. `application/pdf`
    is one of the two types the raw route still serves inline."""
    return Panel(f'<iframe class="view-pdf" src="{_esc(attachment_url)}" title="PDF"></iframe>',
                 kind=kind, label="PDF")


def _spreadsheet(content, filename, content_type, attachment_url, child_url, kind) -> Panel:
    """Every sheet, not the first three — a tracker's delivery tab is rarely the first one."""
    if kind == sniff.KIND_XLS:
        sheets = _read_xls(content)
    else:
        sheets = _read_xlsx(content)
    if not sheets:
        return Panel(_note("This spreadsheet has no readable cells."), kind=kind,
                     label="Spreadsheet")
    blocks = []
    for index, (title, rows) in enumerate(sheets):
        clipped = " · first %d rows" % MAX_TABLE_ROWS if len(rows) >= MAX_TABLE_ROWS else ""
        blocks.append(_note(f"Sheet “{title}”{clipped}") + _table(rows, f"view-sheet-{index}"))
    return Panel("".join(blocks), kind=kind, label="Spreadsheet")


def _read_xlsx(content: bytes):
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    try:
        out = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(max_row=MAX_TABLE_ROWS, max_col=MAX_TABLE_COLS,
                                       values_only=True):
                if any(cell is not None and str(cell).strip() for cell in row):
                    rows.append(list(row))
            if rows:
                out.append((sheet.title, rows))
        return out
    finally:
        workbook.close()


def _read_xls(content: bytes):
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    out = []
    for sheet in book.sheets():
        rows = []
        for index in range(min(sheet.nrows, MAX_TABLE_ROWS)):
            row = sheet.row_values(index)[:MAX_TABLE_COLS]
            if any(str(cell).strip() for cell in row):
                rows.append(row)
        if rows:
            out.append((sheet.name, rows))
    return out


def _document(content, filename, content_type, attachment_url, child_url, kind) -> Panel:
    """Paragraphs **and tables**, in document order.

    The old preview took `document.paragraphs` only, which is exactly the wrong half: python-docx
    does not report table cells as paragraphs, and a delivery confirmation attached as .docx puts
    its quantities in a table. Reading the body's XML children in order is what keeps a heading
    attached to the table it introduces.
    """
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(content))
    blocks, text_run = [], []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            line = Paragraph(child, document).text.strip()
            if line:
                text_run.append(line)
        elif child.tag.endswith("}tbl"):
            if text_run:
                blocks.append(_pre("\n".join(text_run)))
                text_run = []
            rows = [[cell.text for cell in row.cells] for row in Table(child, document).rows]
            if rows:
                blocks.append(_table(rows, f"view-doc-{len(blocks)}"))
    if text_run:
        blocks.append(_pre("\n".join(text_run)))
    return Panel("".join(blocks) or _note("This document is empty."), kind=kind, label="Document")


def _presentation(content, filename, content_type, attachment_url, child_url, kind) -> Panel:
    """Slide by slide. A deck is never mined for delivery records — `unsupported_adapter` says so —
    but someone still has to be able to read one without leaving the browser."""
    from pptx import Presentation

    deck = Presentation(io.BytesIO(content))
    blocks = []
    for number, slide in enumerate(deck.slides, start=1):
        lines, tables = [], []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                tables.append([[cell.text for cell in row.cells] for row in shape.table.rows])
        if not lines and not tables:
            continue
        blocks.append(_note(f"Slide {number}"))
        if lines:
            blocks.append(_pre("\n".join(lines)))
        for index, rows in enumerate(tables):
            blocks.append(_table(rows, f"view-slide-{number}-{index}"))
    return Panel("".join(blocks) or _note("This deck has no readable text."), kind=kind,
                 label="Presentation")


def _html(content, filename, content_type, attachment_url, child_url, kind) -> Panel:
    """The sender's own markup, in a frame that cannot run it.

    This is the attachment the raw route refuses to serve inline, and rightly — served from our
    origin it would be script running as trusted page code. A `srcdoc` sandbox without
    `allow-scripts` is a different thing: the markup renders, and nothing in it executes.
    """
    return Panel(_frame(_decode(content)), kind=kind, label="Web page")


def _text(content, filename, content_type, attachment_url, child_url, kind) -> Panel:
    """Delimited text becomes a table; RTF is de-encapsulated; everything else is shown as it is.

    `looks_delimited` is the extraction cascade's own test, reused so the viewer and the parser
    agree about what counts as a table — a tracker exported as `.txt` is still a tracker.
    """
    text = _decode(content)
    if filename.lower().endswith(".rtf") or text.lstrip().startswith("{\\rtf"):
        return Panel(_pre(_strip_rtf(text)), kind=kind, label="Rich text")

    delimiter = text_adapter.looks_delimited(text, filename or "")
    if delimiter:
        import csv as csv_module

        rows = [row for row in csv_module.reader(io.StringIO(text), delimiter=delimiter)
                if any(cell.strip() for cell in row)][:MAX_TABLE_ROWS]
        if rows:
            return Panel(_table(rows, "view-csv"), kind=kind, label="Delimited text")
    return Panel(_pre(text), kind=kind, label="Text")


def _strip_rtf(text: str) -> str:
    """`striprtf` for a plain RTF document, falling back to the cascade's own stripper.

    `text_adapter.strip_rtf` leads with RTFDE, which is built for the *encapsulated-HTML* case —
    an RTF body carrying an HTML alternative — and returns nothing useful for an ordinary RTF
    attachment. Trying the general reader first and that one second covers both.
    """
    try:
        from striprtf.striprtf import rtf_to_text

        out = rtf_to_text(text, errors="ignore").strip()
        if out:
            return out
    except Exception:                                           # noqa: BLE001
        pass
    return text_adapter.strip_rtf(text)


# --- Containers ---------------------------------------------------------------

def _message(content, filename, content_type, attachment_url, child_url, kind) -> Panel:
    """A `.msg` or `.eml` attachment, read as the message it is.

    This is where Premier's delivery evidence actually lives: the POD date, the signature, the
    carrier and the real received quantity all arrived inside a forwarded Delivered Notification,
    and for a long time that attachment was mis-sniffed and recorded as holding nothing. Rendering
    it as a message — its own header, its own body, its own attachments, each openable — is the
    difference between seeing the evidence and seeing a filename.
    """
    children = _expand(content, filename, content_type)
    header, body_html, body_text = _message_parts(content, filename, content_type)

    blocks = [_headers(header)]
    if body_html:
        blocks.append(_frame(body_html))
    elif body_text:
        blocks.append(_pre(body_text))
    else:
        blocks.append(_note("This message had no body."))
    blocks.append(_children_list(children, child_url, "Attachments"))
    return Panel("".join(blocks), kind=kind, label="Email message",
                 children=_child_refs(children))


def _message_parts(content: bytes, filename: str, content_type: str):
    """`(headers, body_html, body_text)`. Handles both shapes an attached message arrives in.

    Graph returns an `itemAttachment` as RFC-822 MIME even when the filename says `.msg` — that
    mismatch is what once emptied every POD field in thirteen of thirteen records — so the format
    is decided by `sniff.is_eml`, never by the extension.
    """
    if sniff.is_eml(content):
        import email
        from email import policy

        message = email.message_from_bytes(content, policy=policy.default)
        header = [(name, str(message.get(name, ""))) for name in
                  ("From", "To", "Subject", "Date") if message.get(name)]
        html_part = message.get_body(preferencelist=("html",))
        text_part = message.get_body(preferencelist=("plain",))
        return (
            header,
            html_part.get_content() if html_part is not None else "",
            text_part.get_content() if text_part is not None else "",
        )

    import extract_msg

    message = extract_msg.openMsg(io.BytesIO(content))
    try:
        header = [(name, value) for name, value in (
            ("From", message.sender or ""), ("To", message.to or ""),
            ("Subject", message.subject or ""), ("Date", str(message.date or "")),
        ) if value]
        return header, message.htmlBody and _decode(message.htmlBody) or "", message.body or ""
    finally:
        message.close()


def _archive(content, filename, content_type, attachment_url, child_url, kind) -> Panel:
    """What is inside, so someone can decide whether to go and get it.

    Zips are expanded through the same `ZipContainerAdapter` the pipeline uses, under the same
    budget and the same zip-bomb guards — a viewer must not become the one path that will happily
    inflate a 200:1 archive. `.tar`/`.gz` gets a listing the pipeline itself does not do; it is
    stdlib and it is the difference between "archive" and knowing a POD is in there.
    """
    children = _expand(content, filename, content_type)
    if not children and kind == sniff.KIND_ARCHIVE:
        children = _tar_members(content)
    if not children:
        return _unsupported(content, filename, content_type, attachment_url, child_url, kind)
    return Panel(_children_list(children, child_url, "Files inside"), kind=kind, label="Archive",
                 children=_child_refs(children))


def _tar_members(content: bytes):
    import tarfile

    from pipeline.models import Attachment

    try:
        with tarfile.open(fileobj=io.BytesIO(content)) as archive:
            out = []
            for member in archive.getmembers()[:MAX_MEMBERS]:
                if not member.isfile():
                    continue
                oversize = member.size > settings.MAX_ATTACHMENT_BYTES
                data = b"" if oversize else (archive.extractfile(member) or io.BytesIO()).read()
                out.append(Attachment(
                    filename=member.name, content_type="", content_bytes=data,
                    size_bytes=member.size,
                    drop_hint=f"oversize:{member.size} bytes" if oversize else None,
                ))
            return out
    except Exception:                                           # noqa: BLE001
        return []


def _expand(content: bytes, filename: str, content_type: str):
    """Members of a container, expanded now rather than read out of the ledger.

    Deliberately not ledger-driven: every attachment row in both live stores is `depth = 0`, so no
    container has ever actually been expanded there, and a viewer that could only open
    pre-expanded children would open nothing. The adapters, the shared `Budget` and
    `MAX_CONTAINER_DEPTH` are the pipeline's own, so this stays inside the limits ingest respects.
    """
    source = ExtractionSource(
        source_email_id="", email_date="", source_type="attachment",
        filename=filename or "", content_type=content_type or "", content_bytes=content,
    )
    for adapter in (containers.MsgContainerAdapter(), containers.ZipContainerAdapter()):
        if adapter.can_expand(source):
            try:
                return adapter.expand(source, containers.Budget.fresh())
            except Exception:                                   # noqa: BLE001
                return []
    return []


def _child_refs(children) -> List[ChildRef]:
    return [ChildRef(index=index, filename=child.filename or f"part {index + 1}",
                     size_bytes=child.size_bytes or len(child.content_bytes or b""),
                     note=child.drop_hint or "")
            for index, child in enumerate(children)]


def child_at(content: bytes, filename: str, content_type: str, path: Sequence[int]):
    """Walk into a container by index path — `[0, 2]` is the third member of the first member.

    Returns `(bytes, filename, content_type)` or None. Bounded by `MAX_CONTAINER_DEPTH`, the same
    ceiling `dispatch` enforces, so a message nested inside itself cannot be walked for ever.
    """
    current, name, mime = content, filename, content_type
    for depth, index in enumerate(path):
        if depth >= settings.MAX_CONTAINER_DEPTH:
            return None
        children = _expand(current, name, mime)
        if index < 0 or index >= len(children):
            return None
        child = children[index]
        current, name, mime = child.content_bytes or b"", child.filename or "", child.content_type
    return current, name, mime


# --- Everything else ----------------------------------------------------------

def _unsupported(content, filename, content_type, attachment_url, child_url, kind) -> Panel:
    """No reader, and an honest panel rather than a blank one.

    The guidance is `unsupported_adapter.GUIDANCE` — the same sentence the ledger records for this
    kind — so the screen and the audit trail say the same thing. A `.doc` has no reliable
    pure-Python reader and saying so, with what to do about it, is the correct answer; showing
    nothing is not.
    """
    guidance = GUIDANCE.get(kind, GUIDANCE[sniff.KIND_UNKNOWN])
    head = content[:16].hex()
    return Panel(
        _note(guidance)
        + _note(f"First bytes: {head}")
        + _note("Use Download to open it outside the browser."),
        kind=kind, label=_LABELS.get(kind, "File"),
    )


_RENDERERS: Dict[str, Callable] = {
    sniff.KIND_IMAGE: _image,
    sniff.KIND_PDF: _pdf,
    sniff.KIND_XLSX: _spreadsheet,
    sniff.KIND_XLS: _spreadsheet,
    sniff.KIND_DOCX: _document,
    sniff.KIND_PPTX: _presentation,
    sniff.KIND_HTML: _html,
    sniff.KIND_TEXT: _text,
    sniff.KIND_MSG: _message,
    sniff.KIND_ZIP_OFFICE: _archive,
    sniff.KIND_ARCHIVE: _archive,
}
"""Kind → renderer. Everything absent falls to `_unsupported`, which is a panel and not a blank.

`tests/test_attachment_view.py` walks `sniff.ALL_KINDS` and requires every one to produce a
non-empty panel, so a sixteenth kind cannot be added without deciding what shows it."""

_LABELS = {
    sniff.KIND_DOC: "Word 97-2003 document",
    sniff.KIND_LEGACY_OFFICE: "Legacy Office file",
    sniff.KIND_ARCHIVE_UNSUPPORTED: "Archive (7-Zip or RAR)",
    sniff.KIND_ARCHIVE: "Archive",
    sniff.KIND_ZIP_OFFICE: "OOXML archive",
    sniff.KIND_UNKNOWN: "Unrecognised file",
}


# --- Markup helpers -----------------------------------------------------------

def _frame(document: str) -> str:
    """Untrusted markup, rendered and unable to act.

    Identical in posture to `mail_view._body_frame`: no `allow-scripts`, `default-src 'none'`, and
    `allow-same-origin` only so the parent can measure the frame and size it. Both flags together
    would be no sandbox at all — `test_operations_readonly.test_message_frame_never_allows_scripts`
    holds that line across this package.
    """
    csp = ("<meta http-equiv='Content-Security-Policy' content=\"default-src 'none'; "
           "img-src 'self' data:; style-src 'unsafe-inline'; font-src data:\">")
    srcdoc = _esc(f"<!doctype html><html><head><meta charset='utf-8'>{csp}"
                  f"<base target='_blank'></head><body style='margin:12px'>{document}</body></html>")
    return (f'<iframe class="view-frame" '
            f'sandbox="allow-same-origin allow-popups allow-popups-to-escape-sandbox" '
            f'srcdoc="{srcdoc}" title="Attachment"></iframe>')


def _table(rows: Sequence[Sequence], table_id: str) -> str:
    """A grid, paged by the shared client-side pager.

    The pager markup is spelled out rather than built by `api.ui.html.pager()` because this package
    imports neither interface — the same constraint `mail_view` documents, and the same test
    (`test_the_attachment_pager_matches_the_shared_one`) pins the two together.
    """
    width = max((len(row) for row in rows), default=0)
    body = "".join(
        "<tr>" + "".join(f"<td>{_esc(cell)}</td>" for cell in list(row) + [""] * (width - len(row)))
        + "</tr>"
        for row in rows
    )
    return (f'<div class="scroll" id="{_esc(table_id)}">'
            f'<table class="sheet"><tbody>{body}</tbody></table></div>'
            f'<div class="pager" data-pager-for="{_esc(table_id)}" data-page-size="50" '
            f'data-unit="row">'
            f'<button type="button" class="btn ghost small" data-page="prev">‹ Prev</button>'
            f'<span class="pager-label"></span>'
            f'<button type="button" class="btn ghost small" data-page="next">Next ›</button>'
            f'</div>')


def _children_list(children, child_url: str, title: str) -> str:
    """Members of a container, each a control that opens it in this same viewer.

    A member with a `drop_hint` — oversize, corrupt, past the budget — is listed and *not* made
    clickable, with the reason shown. Silently omitting it would be the container equivalent of the
    blank panel this module exists to remove.
    """
    if not children:
        return _note("Nothing inside this one.")
    rows = []
    for index, child in enumerate(children):
        name = child.filename or f"part {index + 1}"
        size = _size(child.size_bytes or len(child.content_bytes or b""))
        if child.drop_hint or not child.content_bytes:
            rows.append(f'<li class="view-child"><span>{_esc(name)}</span>'
                        f'<span class="muted">{_esc(size)} · '
                        f'{_esc(child.drop_hint or "no bytes")}</span></li>')
            continue
        target = f"{child_url}{index}" if child_url else ""
        rows.append(
            f'<li class="view-child"><button type="button" class="link-btn" '
            f'data-frag="{_esc(target)}">{_esc(name)}</button>'
            f'<span class="muted">{_esc(size)}</span></li>'
            if target else
            f'<li class="view-child"><span>{_esc(name)}</span>'
            f'<span class="muted">{_esc(size)}</span></li>'
        )
    return (f'<h4 class="att-title">{_esc(title)} ({len(children)})</h4>'
            f'<ul class="view-children">{"".join(rows)}</ul>')


def _headers(pairs) -> str:
    if not pairs:
        return ""
    rows = "".join(f'<div class="view-hdr-row"><span class="muted">{_esc(name)}</span>'
                   f'<span>{_esc(value)}</span></div>' for name, value in pairs)
    return f'<div class="view-hdr">{rows}</div>'


def _decode(content) -> str:
    if isinstance(content, str):
        return content
    for encoding in ("utf-8", "utf-16", "cp1252"):
        try:
            return content.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return content.decode("latin-1", "replace")


def _pre(text: str) -> str:
    clipped = text[:MAX_TEXT_CHARS]
    tail = _note(f"Showing the first {MAX_TEXT_CHARS:,} characters.") if len(text) > MAX_TEXT_CHARS \
        else ""
    return f'<pre class="att-text">{_esc(clipped)}</pre>{tail}'


def _note(text: str) -> str:
    return f'<p class="note">{_esc(text)}</p>'


def _size(n: int) -> str:
    if n < 1024:
        return f"{n} B"
    if n < 1024 * 1024:
        return f"{n / 1024:.0f} KB"
    return f"{n / 1024 / 1024:.1f} MB"


def _esc(value) -> str:
    """Escapes anything, not just `str` — spreadsheet cells arrive as ints, floats, datetimes and
    None. Same reason and same semantics as `mail_view._esc`."""
    if value is None:
        return ""
    return escape(value if isinstance(value, str) else str(value))
